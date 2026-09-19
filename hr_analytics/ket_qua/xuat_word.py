"""Xuất báo cáo PeopleRisk AI ra file Word (.docx)."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


def _safe(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).replace("\r\n", "\n").replace("\r", "\n")
    return "".join(ch if ord(ch) >= 32 or ch == "\n" else " " for ch in text)


def _add_heading(doc: Document, title: str, level: int = 1) -> None:
    heading = doc.add_heading(level=level)
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
    heading_run = heading.runs[0] if heading.runs else heading.add_run(title)
    heading_run.text = title
    heading_run.bold = True
    heading_run.font.size = Pt(14 if level == 1 else 12)


def _add_paragraph(doc: Document, text: str, *, bold: bool = False, italic: bool = False, size: int = 11) -> None:
    p = doc.add_paragraph()
    run = p.add_run(_safe(text))
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)


def _add_bullet(doc: Document, text: str, *, size: int = 10) -> None:
    p = doc.add_paragraph(style="List Bullet")
    run = p.add_run(_safe(text))
    run.font.size = Pt(size)


def _add_kv(doc: Document, label: str, value: Any) -> None:
    p = doc.add_paragraph()
    label_run = p.add_run(f"{_safe(label)}: ")
    label_run.bold = True
    label_run.font.size = Pt(11)
    value_run = p.add_run(_safe(value))
    value_run.font.size = Pt(11)


def _add_table_from_mapping(doc: Document, rows: list[tuple[str, Any]]) -> None:
    if not rows:
        _add_paragraph(doc, "Không có dữ liệu.")
        return
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    header = table.rows[0].cells
    header[0].text = "Chỉ số"
    header[1].text = "Giá trị"
    for cell in header:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
    for label, value in rows:
        row_cells = table.add_row().cells
        row_cells[0].text = _safe(label)
        row_cells[1].text = _safe(value)


def _format_metric_value(value: Any) -> str:
    if value is None:
        return "—"
    try:
        if isinstance(value, (int, float)):
            if abs(float(value)) >= 1000:
                return f"{float(value):,.0f}"
            return f"{float(value):,.4f}" if float(value) < 1 else f"{float(value):,.2f}"
    except (TypeError, ValueError):
        pass
    return str(value)


def export_report_word(
    path: str | Path,
    *,
    dataset_name: str | None,
    target: str | None,
    n_rows: int,
    n_cols: int,
    filter_note: str,
    kpis: dict[str, Any],
    insights: list[dict[str, Any]],
    recommendations: list[dict[str, Any]],
    model_name: str | None = None,
    model_metrics: dict[str, Any] | None = None,
    metrics_table: Any = None,
    generated_at: datetime | None = None,
) -> Path:
    """Tạo file Word báo cáo tổng hợp. Trả về đường dẫn đã lưu."""
    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    when = generated_at or datetime.now()

    doc = Document()
    title = doc.add_paragraph()
    title_run = title.add_run("PeopleRisk AI — Báo cáo phân tích nhân sự")
    title_run.bold = True
    title_run.font.size = Pt(20)

    _add_paragraph(doc, f"Thời điểm xuất: {when.strftime('%d/%m/%Y %H:%M:%S')}")
    _add_paragraph(doc, f"Dataset: {dataset_name or '—'}")
    _add_paragraph(doc, f"Kích thước: {n_rows:,} dòng × {n_cols} cột".replace(",", "."))
    _add_paragraph(doc, f"Target: {target or '—'}")
    _add_paragraph(doc, f"Bộ lọc áp dụng: {filter_note or 'Không lọc'}")
    if model_name:
        _add_paragraph(doc, f"Mô hình đang dùng: {model_name}")

    _add_heading(doc, "1. Thông tin báo cáo", level=1)
    _add_kv(doc, "Thời điểm xuất", when.strftime("%d/%m/%Y %H:%M:%S"))
    _add_kv(doc, "Dataset", dataset_name or "—")
    _add_kv(doc, "Kích thước", f"{n_rows:,} dòng × {n_cols} cột".replace(",", "."))
    _add_kv(doc, "Target", target or "—")
    _add_kv(doc, "Bộ lọc đang áp dụng", filter_note or "Không lọc")
    if model_name:
        _add_kv(doc, "Mô hình đang dùng", model_name)

    _add_heading(doc, "2. Chỉ số tổng quan (KPI)", level=1)
    mapping = [
        ("Tổng nhân sự", kpis.get("total_employees")),
        ("Đang làm việc", kpis.get("employees_staying")),
        ("Đã nghỉ việc", kpis.get("employees_leaving")),
        ("Tỷ lệ nghỉ việc (%)", kpis.get("attrition_rate")),
        ("Tuổi trung bình", kpis.get("avg_age")),
        ("Thu nhập trung bình", kpis.get("avg_income")),
        ("Thâm niên TB (năm)", kpis.get("avg_years_company")),
        ("Hài lòng công việc TB", kpis.get("avg_job_satisfaction")),
    ]
    valid = [(label, val) for label, val in mapping if val is not None]
    _add_table_from_mapping(doc, valid)

    _add_heading(doc, "3. Kết quả mô hình", level=1)
    if model_metrics:
        _add_paragraph(doc, f"Model được chọn: {model_name or '—'}")
        metric_rows = []
        for key in ["Accuracy", "Precision", "Recall", "F1", "ROC-AUC"]:
            if key in model_metrics:
                try:
                    metric_rows.append((key, f"{float(model_metrics[key]):.4f}"))
                except (TypeError, ValueError):
                    metric_rows.append((key, model_metrics[key]))
        if metric_rows:
            _add_table_from_mapping(doc, metric_rows)
        else:
            _add_paragraph(doc, "Không có chỉ số mô hình.")
    elif metrics_table is not None and hasattr(metrics_table, "empty") and not metrics_table.empty:
        _add_paragraph(doc, "Bảng so sánh mô hình:")
        cols = [c for c in metrics_table.columns if c in ("Model", "Accuracy", "Precision", "Recall", "F1", "ROC-AUC")]
        if not cols:
            cols = list(metrics_table.columns[:6])
        table = doc.add_table(rows=1, cols=len(cols))
        table.style = "Table Grid"
        for idx, col in enumerate(cols):
            table.rows[0].cells[idx].text = str(col)
            for paragraph in table.rows[0].cells[idx].paragraphs:
                for run in paragraph.runs:
                    run.bold = True
        for _, row in metrics_table.iterrows():
            r = table.add_row().cells
            for idx, col in enumerate(cols):
                if col in row.index:
                    r[idx].text = _safe(row[col])
                else:
                    r[idx].text = ""
    else:
        _add_paragraph(doc, "Chưa huấn luyện mô hình — phần này trống.")

    _add_heading(doc, f"4. Insights ({len(insights)})", level=1)
    if not insights:
        _add_paragraph(doc, "Không có insight.")
    else:
        for i, ins in enumerate(insights, 1):
            title = ins.get("title") or ins.get("group") or f"Insight {i}"
            severity = ins.get("severity") or ""
            _add_paragraph(doc, f"{i}. {title}{f'  [{severity}]' if severity else ''}", bold=True)
            if ins.get("insight"):
                _add_paragraph(doc, ins["insight"])
            if ins.get("evidence"):
                _add_paragraph(doc, f"Bằng chứng: {ins['evidence']}")

    _add_heading(doc, f"5. Khuyến nghị ({len(recommendations)})", level=1)
    if not recommendations:
        _add_paragraph(doc, "Không có khuyến nghị.")
    else:
        for i, rec in enumerate(recommendations, 1):
            based = rec.get("based_on") or rec.get("problem") or f"Khuyến nghị {i}"
            priority = rec.get("priority") or ""
            _add_paragraph(doc, f"{i}. {based}{f'  [{priority}]' if priority else ''}", bold=True)
            if rec.get("problem"):
                _add_paragraph(doc, f"Vấn đề: {rec['problem']}")
            if rec.get("evidence"):
                _add_paragraph(doc, f"Bằng chứng: {rec['evidence']}")
            if rec.get("recommended_action"):
                _add_paragraph(doc, f"Hành động: {rec['recommended_action']}")
            if rec.get("expected_goal"):
                _add_paragraph(doc, f"Mục tiêu: {rec['expected_goal']}")

    _add_heading(doc, "6. Kết luận tóm tắt", level=1)
    rate = kpis.get("attrition_rate")
    summary_lines = [
        f"Dataset «{dataset_name or '—'}» với {n_rows:,} nhân viên đang phân tích (target: {target or '—'}).".replace(",", "."),
    ]
    if rate is not None:
        summary_lines.append(f"Tỷ lệ nghỉ việc hiện tại khoảng {rate}%.")
    if model_name:
        summary_lines.append(f"Mô hình dự báo đang dùng: {model_name}.")
    if insights:
        summary_lines.append(f"Đã rút {len(insights)} insight và {len(recommendations)} khuyến nghị hành động.")
    else:
        summary_lines.append("Chưa có đủ insight — kiểm tra lại dữ liệu hoặc bộ lọc.")
    summary_lines.append(
        "Báo cáo được sinh tự động từ dữ liệu và bộ lọc đang áp dụng trong PeopleRisk AI; không thay thế phán đoán chuyên môn HR."
    )
    for line in summary_lines:
        _add_bullet(doc, line)

    doc.save(str(out))
    return out
